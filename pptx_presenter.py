#!/usr/bin/env python3
"""
PowerPoint Presenter - View and present .pptx files without PowerPoint.
Files stay in .pptx format for sharing with PowerPoint users.

Usage:
    python3 pptx_presenter.py [your_file.pptx]
    Then open http://localhost:5050 in your browser.
    If no file is given, you can upload one through the browser.
"""

import sys
import os
import io
import base64
import json
from pathlib import Path

from flask import Flask, render_template_string, request, redirect, url_for, jsonify, send_from_directory
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from PIL import Image

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'uploads')
app.config['MAX_CONTENT_LENGTH'] = 200 * 1024 * 1024  # 200MB max

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

CURRENT_FILE = None


def rgb_to_css(rgb_color):
    if rgb_color is None:
        return None
    try:
        return f"rgb({rgb_color[0]}, {rgb_color[1]}, {rgb_color[2]})"
    except (TypeError, IndexError):
        return None


def get_shape_color(shape):
    try:
        fill = shape.fill
        if fill.type is not None:
            if fill.type == 1:  # solid
                return rgb_to_css(fill.fore_color.rgb)
    except Exception:
        pass
    return None


def extract_slide_data(pptx_path):
    prs = Presentation(pptx_path)
    slide_width = prs.slide_width
    slide_height = prs.slide_height
    w_px = 960
    h_px = int(w_px * slide_height / slide_width)
    scale = w_px / slide_width

    slides_data = []

    for slide_num, slide in enumerate(prs.slides, 1):
        slide_info = {
            "number": slide_num,
            "elements": [],
            "background": "#FFFFFF",
            "width": w_px,
            "height": h_px,
            "notes": ""
        }

        # Extract background
        try:
            bg = slide.background
            if bg.fill.type is not None and bg.fill.type == 1:
                slide_info["background"] = rgb_to_css(bg.fill.fore_color.rgb) or "#FFFFFF"
        except Exception:
            pass

        # Extract notes
        try:
            if slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                slide_info["notes"] = notes_frame.text
        except Exception:
            pass

        for shape in slide.shapes:
            elem = {
                "left": int(shape.left * scale) if shape.left else 0,
                "top": int(shape.top * scale) if shape.top else 0,
                "width": int(shape.width * scale) if shape.width else 0,
                "height": int(shape.height * scale) if shape.height else 0,
                "type": "unknown",
            }

            # Images
            if shape.shape_type == 13:  # Picture
                try:
                    image = shape.image
                    img_bytes = image.blob
                    content_type = image.content_type
                    b64 = base64.b64encode(img_bytes).decode('utf-8')
                    elem["type"] = "image"
                    elem["src"] = f"data:{content_type};base64,{b64}"
                    slide_info["elements"].append(elem)
                except Exception:
                    pass
                continue

            # Shape background
            bg_color = get_shape_color(shape)
            if bg_color:
                elem["background"] = bg_color

            # Tables
            if shape.has_table:
                table = shape.table
                rows = []
                for row in table.rows:
                    cells = []
                    for cell in row.cells:
                        cells.append(cell.text)
                    rows.append(cells)
                elem["type"] = "table"
                elem["rows"] = rows
                slide_info["elements"].append(elem)
                continue

            # Text
            if shape.has_text_frame:
                paragraphs = []
                for para in shape.text_frame.paragraphs:
                    p_data = {"runs": [], "alignment": "left", "level": para.level or 0}

                    if para.alignment == PP_ALIGN.CENTER:
                        p_data["alignment"] = "center"
                    elif para.alignment == PP_ALIGN.RIGHT:
                        p_data["alignment"] = "right"

                    for run in para.runs:
                        run_data = {"text": run.text}
                        font = run.font

                        if font.size:
                            run_data["size"] = int(font.size.pt * scale * 1.1)
                        if font.bold:
                            run_data["bold"] = True
                        if font.italic:
                            run_data["italic"] = True
                        if font.underline:
                            run_data["underline"] = True
                        if font.color and font.color.rgb:
                            run_data["color"] = rgb_to_css(font.color.rgb)
                        if font.name:
                            run_data["font"] = font.name

                        p_data["runs"].append(run_data)

                    if not p_data["runs"] and para.text:
                        p_data["runs"].append({"text": para.text})

                    paragraphs.append(p_data)

                if any(r["text"].strip() for p in paragraphs for r in p["runs"]):
                    elem["type"] = "text"
                    elem["paragraphs"] = paragraphs
                    slide_info["elements"].append(elem)

        slides_data.append(slide_info)

    return slides_data, len(prs.slides)


HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{{ title }} - PowerPoint Presenter</title>
<style>
* { margin: 0; padding: 0; box-sizing: border-box; }
body {
    font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
    background: #1a1a2e;
    color: #eee;
    overflow: hidden;
    height: 100vh;
}

/* Top bar */
.toolbar {
    display: flex;
    align-items: center;
    justify-content: space-between;
    padding: 8px 20px;
    background: #16213e;
    border-bottom: 1px solid #0f3460;
    height: 50px;
    z-index: 100;
}
.toolbar .title { font-size: 14px; color: #a0a0c0; }
.toolbar .controls { display: flex; gap: 10px; align-items: center; }
.toolbar button {
    background: #0f3460;
    border: 1px solid #533483;
    color: #e0e0ff;
    padding: 6px 16px;
    border-radius: 4px;
    cursor: pointer;
    font-size: 13px;
}
.toolbar button:hover { background: #533483; }
.slide-counter { font-size: 14px; color: #a0a0c0; margin: 0 10px; }

/* Main area */
.main-area {
    display: flex;
    height: calc(100vh - 50px);
}

/* Slide panel (thumbnails) */
.slide-panel {
    width: 180px;
    background: #16213e;
    overflow-y: auto;
    padding: 10px;
    border-right: 1px solid #0f3460;
}
.slide-panel.hidden { display: none; }
.thumb {
    border: 2px solid transparent;
    border-radius: 4px;
    margin-bottom: 8px;
    cursor: pointer;
    padding: 4px;
    transition: border-color 0.2s;
}
.thumb.active { border-color: #e94560; }
.thumb:hover { border-color: #533483; }
.thumb-label { font-size: 11px; text-align: center; color: #888; margin-top: 2px; }
.thumb-box {
    background: #fff;
    border-radius: 2px;
    height: 90px;
    display: flex;
    align-items: center;
    justify-content: center;
    font-size: 24px;
    color: #333;
    font-weight: bold;
}

/* Slide view */
.slide-view {
    flex: 1;
    display: flex;
    align-items: center;
    justify-content: center;
    padding: 20px;
    position: relative;
}

.slide-container {
    position: relative;
    box-shadow: 0 4px 30px rgba(0,0,0,0.5);
    border-radius: 4px;
    overflow: hidden;
    max-width: 100%;
    max-height: 100%;
}

.slide-element {
    position: absolute;
    overflow: hidden;
}

.slide-element.text-element {
    overflow: visible;
}

.slide-element img {
    width: 100%;
    height: 100%;
    object-fit: contain;
}

.slide-element table {
    width: 100%;
    height: 100%;
    border-collapse: collapse;
    font-size: 12px;
    color: #333;
}
.slide-element td {
    border: 1px solid #ccc;
    padding: 4px 8px;
}

.para { margin: 2px 0; }
.para.level-1 { margin-left: 20px; }
.para.level-2 { margin-left: 40px; }

/* Nav arrows */
.nav-arrow {
    position: absolute;
    top: 50%;
    transform: translateY(-50%);
    background: rgba(15, 52, 96, 0.7);
    border: none;
    color: #e0e0ff;
    font-size: 28px;
    padding: 15px 10px;
    cursor: pointer;
    border-radius: 4px;
    z-index: 10;
}
.nav-arrow:hover { background: rgba(83, 52, 131, 0.9); }
.nav-arrow.left { left: 10px; }
.nav-arrow.right { right: 10px; }

/* Notes panel */
.notes-panel {
    width: 280px;
    background: #16213e;
    border-left: 1px solid #0f3460;
    padding: 15px;
    overflow-y: auto;
}
.notes-panel.hidden { display: none; }
.notes-panel h3 { font-size: 13px; color: #e94560; margin-bottom: 8px; }
.notes-panel p { font-size: 13px; line-height: 1.5; color: #c0c0d0; white-space: pre-wrap; }

/* Fullscreen mode */
body.fullscreen .toolbar { display: none; }
body.fullscreen .slide-panel { display: none; }
body.fullscreen .notes-panel { display: none; }
body.fullscreen .main-area { height: 100vh; }
body.fullscreen .slide-view { background: #000; }
body.fullscreen .nav-arrow { opacity: 0; transition: opacity 0.3s; }
body.fullscreen .slide-view:hover .nav-arrow { opacity: 1; }
</style>
</head>
<body>

<div class="toolbar">
    <div class="title">{{ title }}</div>
    <div class="controls">
        <button onclick="togglePanel()">Thumbnails</button>
        <button onclick="toggleNotes()">Notes</button>
        <button class="nav-arrow-btn" onclick="prevSlide()">&#9664;</button>
        <span class="slide-counter"><span id="currentNum">1</span> / {{ total }}</span>
        <button class="nav-arrow-btn" onclick="nextSlide()">&#9654;</button>
        <button onclick="toggleFullscreen()">Present (F5)</button>
    </div>
</div>

<div class="main-area">
    <div class="slide-panel" id="slidePanel">
        {% for s in slides %}
        <div class="thumb {% if loop.first %}active{% endif %}" onclick="goToSlide({{ loop.index0 }})" id="thumb-{{ loop.index0 }}">
            <div class="thumb-box" style="background: {{ s.background }};">{{ loop.index }}</div>
            <div class="thumb-label">Slide {{ loop.index }}</div>
        </div>
        {% endfor %}
    </div>

    <div class="slide-view">
        <button class="nav-arrow left" onclick="prevSlide()">&#9664;</button>

        {% for s in slides %}
        <div class="slide-container" id="slide-{{ loop.index0 }}"
             style="width: {{ s.width }}px; height: {{ s.height }}px; background: {{ s.background }};
                    display: {% if loop.first %}block{% else %}none{% endif %};">
            {% for elem in s.elements %}
                {% if elem.type == 'image' %}
                <div class="slide-element" style="left:{{ elem.left }}px; top:{{ elem.top }}px; width:{{ elem.width }}px; height:{{ elem.height }}px;">
                    <img src="{{ elem.src }}" alt="image">
                </div>
                {% elif elem.type == 'table' %}
                <div class="slide-element" style="left:{{ elem.left }}px; top:{{ elem.top }}px; width:{{ elem.width }}px; height:{{ elem.height }}px;">
                    <table>
                    {% for row in elem.rows %}
                        <tr>{% for cell in row %}<td>{{ cell }}</td>{% endfor %}</tr>
                    {% endfor %}
                    </table>
                </div>
                {% elif elem.type == 'text' %}
                <div class="slide-element text-element" style="left:{{ elem.left }}px; top:{{ elem.top }}px; width:{{ elem.width }}px; height:{{ elem.height }}px;
                    {% if elem.background %}background: {{ elem.background }};{% endif %}">
                    {% for p in elem.paragraphs %}
                    <div class="para level-{{ p.level }}" style="text-align: {{ p.alignment }};">
                        {% for r in p.runs %}
                        <span style="
                            {% if r.size %}font-size: {{ r.size }}px;{% endif %}
                            {% if r.bold %}font-weight: bold;{% endif %}
                            {% if r.italic %}font-style: italic;{% endif %}
                            {% if r.underline %}text-decoration: underline;{% endif %}
                            {% if r.color %}color: {{ r.color }};{% else %}color: #333;{% endif %}
                            {% if r.font %}font-family: '{{ r.font }}', sans-serif;{% endif %}
                        ">{{ r.text }}</span>
                        {% endfor %}
                    </div>
                    {% endfor %}
                </div>
                {% endif %}
            {% endfor %}
        </div>
        {% endfor %}

        <button class="nav-arrow right" onclick="nextSlide()">&#9654;</button>
    </div>

    <div class="notes-panel" id="notesPanel">
        <h3>Speaker Notes</h3>
        <p id="notesText">{{ slides[0].notes if slides else '' }}</p>
    </div>
</div>

<script>
const totalSlides = {{ total }};
const notes = {{ notes_json|safe }};
let current = 0;

function showSlide(n) {
    document.getElementById('slide-' + current).style.display = 'none';
    document.getElementById('thumb-' + current).classList.remove('active');
    current = n;
    document.getElementById('slide-' + current).style.display = 'block';
    document.getElementById('thumb-' + current).classList.add('active');
    document.getElementById('currentNum').textContent = current + 1;
    document.getElementById('notesText').textContent = notes[current] || '';
    document.getElementById('thumb-' + current).scrollIntoView({block: 'nearest'});
}

function nextSlide() { if (current < totalSlides - 1) showSlide(current + 1); }
function prevSlide() { if (current > 0) showSlide(current - 1); }
function goToSlide(n) { showSlide(n); }

function toggleFullscreen() {
    document.body.classList.toggle('fullscreen');
    if (document.body.classList.contains('fullscreen')) {
        document.documentElement.requestFullscreen?.();
    } else {
        document.exitFullscreen?.();
    }
}

function togglePanel() {
    document.getElementById('slidePanel').classList.toggle('hidden');
}

function toggleNotes() {
    document.getElementById('notesPanel').classList.toggle('hidden');
}

document.addEventListener('keydown', (e) => {
    if (e.key === 'ArrowRight' || e.key === 'ArrowDown' || e.key === ' ' || e.key === 'PageDown') {
        e.preventDefault(); nextSlide();
    } else if (e.key === 'ArrowLeft' || e.key === 'ArrowUp' || e.key === 'PageUp') {
        e.preventDefault(); prevSlide();
    } else if (e.key === 'Home') {
        e.preventDefault(); showSlide(0);
    } else if (e.key === 'End') {
        e.preventDefault(); showSlide(totalSlides - 1);
    } else if (e.key === 'F5' || (e.key === 'F5' && e.ctrlKey)) {
        e.preventDefault(); toggleFullscreen();
    } else if (e.key === 'Escape') {
        document.body.classList.remove('fullscreen');
    }
});
</script>
</body>
</html>
"""

UPLOAD_TEMPLATE = """
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>PowerPoint Presenter - Upload</title>
<style>
* { margin: 0; padding: 0; box-sizing: border-box; }
body {
    font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
    background: #1a1a2e;
    color: #eee;
    display: flex;
    align-items: center;
    justify-content: center;
    height: 100vh;
}
.upload-box {
    background: #16213e;
    border: 2px dashed #533483;
    border-radius: 12px;
    padding: 60px 80px;
    text-align: center;
    cursor: pointer;
    transition: border-color 0.3s;
}
.upload-box:hover { border-color: #e94560; }
.upload-box h1 { font-size: 28px; margin-bottom: 10px; color: #e94560; }
.upload-box p { font-size: 16px; color: #a0a0c0; margin-bottom: 20px; }
.upload-box input[type="file"] { display: none; }
.upload-btn {
    background: #e94560;
    color: #fff;
    border: none;
    padding: 12px 32px;
    font-size: 16px;
    border-radius: 6px;
    cursor: pointer;
}
.upload-btn:hover { background: #c0392b; }
.info { margin-top: 20px; font-size: 13px; color: #666; }
</style>
</head>
<body>
<form class="upload-box" method="POST" enctype="multipart/form-data" action="/upload" id="uploadForm">
    <h1>PowerPoint Presenter</h1>
    <p>Open .pptx files without PowerPoint</p>
    <p style="font-size: 13px;">Your file stays in .pptx format - share it with PowerPoint users anytime</p>
    <input type="file" name="file" id="fileInput" accept=".pptx" onchange="document.getElementById('uploadForm').submit()">
    <button type="button" class="upload-btn" onclick="document.getElementById('fileInput').click()">
        Choose .pptx File
    </button>
    <div class="info">or drag and drop a .pptx file here</div>
</form>
<script>
const box = document.querySelector('.upload-box');
box.addEventListener('dragover', e => { e.preventDefault(); box.style.borderColor = '#e94560'; });
box.addEventListener('dragleave', () => { box.style.borderColor = '#533483'; });
box.addEventListener('drop', e => {
    e.preventDefault();
    const files = e.dataTransfer.files;
    if (files.length && files[0].name.endsWith('.pptx')) {
        document.getElementById('fileInput').files = files;
        document.getElementById('uploadForm').submit();
    }
});
</script>
</body>
</html>
"""


@app.route('/')
def index():
    global CURRENT_FILE
    if CURRENT_FILE and os.path.exists(CURRENT_FILE):
        return redirect(url_for('present'))
    return render_template_string(UPLOAD_TEMPLATE)


@app.route('/upload', methods=['POST'])
def upload():
    global CURRENT_FILE
    if 'file' not in request.files:
        return redirect(url_for('index'))
    file = request.files['file']
    if file.filename == '' or not file.filename.endswith('.pptx'):
        return redirect(url_for('index'))
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
    file.save(filepath)
    CURRENT_FILE = filepath
    return redirect(url_for('present'))


@app.route('/present')
def present():
    global CURRENT_FILE
    if not CURRENT_FILE or not os.path.exists(CURRENT_FILE):
        return redirect(url_for('index'))

    slides_data, total = extract_slide_data(CURRENT_FILE)
    notes = [s["notes"] for s in slides_data]
    title = Path(CURRENT_FILE).stem

    return render_template_string(
        HTML_TEMPLATE,
        slides=slides_data,
        total=total,
        title=title,
        notes_json=json.dumps(notes)
    )


@app.route('/new')
def new_file():
    global CURRENT_FILE
    CURRENT_FILE = None
    return redirect(url_for('index'))


if __name__ == '__main__':
    if len(sys.argv) > 1:
        filepath = sys.argv[1]
        if os.path.exists(filepath) and filepath.endswith('.pptx'):
            CURRENT_FILE = os.path.abspath(filepath)
            print(f"Loaded: {CURRENT_FILE}")
        else:
            print(f"Error: '{filepath}' not found or not a .pptx file")
            sys.exit(1)

    print("\n  PowerPoint Presenter")
    print("  --------------------")
    print("  Open http://localhost:5050 in your browser")
    print("  Press Ctrl+C to stop\n")
    print("  Controls:")
    print("    Arrow keys / Space  - Navigate slides")
    print("    F5                  - Fullscreen presentation")
    print("    Esc                 - Exit fullscreen\n")

    app.run(host='0.0.0.0', port=5050, debug=False)
